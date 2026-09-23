"""Human-readable export formats."""
from pathlib import Path
import json,csv
from docx import Document
from openpyxl import Workbook
from .refi import export_qdpx, export_qdc

def export_all(project_dir,project,codings,codes,sources,audit):
    out=Path(project_dir)/'exports'; out.mkdir(parents=True,exist_ok=True)
    (out/'coding_matrix.json').write_text(json.dumps(codings,ensure_ascii=False,indent=2),encoding='utf-8')
    with (out/'coding_matrix.csv').open('w',newline='',encoding='utf-8') as f:
        fields=['source_name','segment_id','start','end','speaker','quote','code_name','confidence','rationale','status','model']; w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows([{k:x.get(k,'') for k in fields} for x in codings])
    wb=Workbook(); ws=wb.active; ws.title='Coding Matrix'; ws.append(['Source','Segment','Start','End','Speaker','Quote','Code','Confidence','Rationale','Status','Model'])
    for x in codings: ws.append([x.get(k,'') for k in ['source_name','segment_id','start','end','speaker','quote','code_name','confidence','rationale','status','model']])
    wb.save(out/'coding_matrix.xlsx')
    doc=Document(); doc.add_heading(project['title'],0); doc.add_paragraph('ZedThema research audit and coding export'); doc.add_heading('Study design',1)
    for label,key in [('Description','description'),('Research question','research_question'),('Framework','framework'),('Method','method')]: doc.add_paragraph(f'{label}: {project.get(key,"")}')
    doc.add_heading('Coding',1)
    for x in codings: doc.add_paragraph(f"[{x.get('status','')}] {x.get('code_name','')} | {x.get('source_name','')} | {x.get('confidence','')}"); doc.add_paragraph(x.get('quote','')); doc.add_paragraph(x.get('rationale',''))
    doc.save(out/'qualicoder_report.docx')
    (out/'audit_trail.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    try:
        export_qdpx(project_dir,project,codes,sources,codings)
        export_qdc(project_dir,project,codes)
    except Exception as exc: (out/'REFI-QDA-export-error.txt').write_text(str(exc),encoding='utf-8')
    return out
